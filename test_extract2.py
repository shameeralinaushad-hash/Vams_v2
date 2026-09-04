import openpyxl
from io import BytesIO

with open(r'C:\Users\haide\Downloads\Technical_Offer_Annexures_A_B_C (1).xlsx', 'rb') as f:
    data = f.read()

wb = openpyxl.load_workbook(BytesIO(data), data_only=True)

for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n=== {name} ===")
    table = []
    for row in ws.iter_rows(values_only=True):
        row_vals = [str(cell).strip() if cell is not None else "" for cell in row]
        if any(v for v in row_vals if v):
            table.append(row_vals)
    
    # Show the raw table rows
    for i, row in enumerate(table[:12]):
        print(f"  raw_row {i}: {row}")

    # Now check what _find_header_row would return
    from ml.tender_extractor import _find_header_row
    header_idx, header = _find_header_row(table)
    print(f"  Header row index: {header_idx}")
    print(f"  Header: {header}")
    
    if header_idx is not None:
        from ml.tender_extractor import _find_column
        unit_col = _find_column(header, ["unit", "uom", "measure", "pkg", "pack", "per", "a/u"])
        print(f"  Unit col index: {unit_col}")
        if unit_col is not None:
            # Show unit values from data rows
            for i, row in enumerate(table[header_idx+1:header_idx+5]):
                if unit_col < len(row):
                    print(f"    row {header_idx+1+i}, unit_col={unit_col}: '{row[unit_col]}'")
